from libs.Functions import *
import numpy as np
import pandas as pd
from openpyxl.utils import column_index_from_string
def Mail_Merged():
    import openpyxl, docx, shutil
    from os.path import join
###. Nhập liệu
#####.Nhập đường dẫn file
    wb_path = input("Hãy nhập đường dẫn file excel: ")
    wb_path = format_path(wb_path)
#####. Chọn Sheet
    print("File excel gồm các sheet: ")
    dict_sheet = {}
    count = 1
    for sheet in openpyxl.load_workbook(wb_path).sheetnames:
        print(str(count) + ".", sheet)
        dict_sheet[count] = sheet
        count +=1
    option = int(input("Hãy chọn sheet sử dụng: "))
    ws_name = dict_sheet[option]
#####. Chọn các cột sẽ định dạng số
    cot_so = (input("Hãy nhập các cột bạn muốn định dạng số cách nhau bằng dấu \",\": "))
    cot_so_idx = []
    if len(cot_so)>0:
        if " " in cot_so:
            cot_so = cot_so.replace(" ","")
        cot_so = cot_so.split(",")
        for char in cot_so:
            cot_so_idx.append(column_index_from_string(char) - 1)
#####. Chọn các cột sẽ chứa tên file xuất
    cot_ten = (input("Hãy nhập cột chưa tên file xuất: "))
    ten_cl_idx = column_index_from_string(cot_ten) - 1
#####. Nhập đương dẫn Doc sample
    doc_sample_path = input("Hãy nhập đường link file word: ")
    doc_sample_path = format_path(doc_sample_path)
#####. Nhập đương dẫn file xuất
    des_folder = input("Hãy nhập đường link folder xuất file: ")
    des_folder = format_path(des_folder)
#1. INPUT dư liệu excel
    ws = pd.read_excel(wb_path, header = None, sheet_name = ws_name)
#2. Tag «» Mail_merge cho dòng tên cột (dòng 1)
    for cl_idx in range(0, ws.shape[1]):
        ws.iloc[0, cl_idx] = "«" + str(ws.iloc[0, cl_idx]) + "»"
    dict_cl_name = {}
    count = 0
    for name in ws.iloc[0]:
        dict_cl_name[name] = count
        count +=1 
#####. Định dạng lại các cột số được chọn:
    if len(cot_so_idx)>0:
        for idx in cot_so_idx:
            for r_idx in range(1, ws.shape[0]):
                num = str(ws.iloc[r_idx, idx])
                if check_so(num):
                    ws.iloc[r_idx, idx] = num_format(num)
#3. Duyệt từng dòng của file excel
    for r_idx in range(1, ws.shape[0]):
        if str(ws.iloc[r_idx, ten_cl_idx]) != "nan":
#4. Tạo và đặt tên file word kết quả
####. Chỉnh tên muốn đặt
            out_name = str(ws.iloc[r_idx, ten_cl_idx]) + ".docx"
            shutil.copy(doc_sample_path, join(des_folder, out_name))
    #5. Mở từng file vừa tạo để gắn giá trị theo Tag Mail_Merge
            doc = docx.Document(join(des_folder, out_name))
    ####. Duyệt text
            for para in doc.paragraphs:
                for run in para.runs:
                    if run.text in dict_cl_name.keys():
                        content = str(ws.iloc[r_idx, dict_cl_name[run.text]])
                        if content != "nan":
                            run.text = content
                        else:
                            run.text = ""
    ####. Duyệt Table
            for tab in doc.tables:
                for row in tab.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                if run.text in dict_cl_name.keys():
                                    content = str(ws.iloc[r_idx, dict_cl_name[run.text]])
                                    if content != "nan":
                                        run.text = content
                                    else:
                                        run.text = "" 
    ####. Duyệt header và footer
            for sec in doc.sections:
                for para in sec.header.paragraphs:
                    for run in para.runs:
                        if run.text in dict_cl_name.keys():
                            content = str(ws.iloc[r_idx, dict_cl_name[run.text]])
                            if content != "nan":
                                run.text = content
                            else:
                                run.text = ""
                for para in sec.footer.paragraphs:
                    for run in para.runs:
                        if run.text in dict_cl_name.keys():
                            content = str(ws.iloc[r_idx, dict_cl_name[run.text]])
                            if content != "nan":
                                run.text = content
                            else:
                                run.text = "" 
            doc.save(join(des_folder,out_name))